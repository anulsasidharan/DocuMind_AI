"""
Generate docs/GCP_deployment_guide.docx — a complete step-by-step Terraform
deployment guide for DocuMind AI on Google Cloud Platform.

Run from repo root:
    python scripts/generate_deployment_guide.py
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys

OUT = Path(__file__).parent.parent / "docs" / "GCP_deployment_guide.docx"

# ── helpers ──────────────────────────────────────────────────────────────────

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def add_code(doc, code: str):
    """Add a monospaced code block paragraph."""
    p = doc.add_paragraph()
    p.style = doc.styles["No Spacing"]
    run = p.add_run(code)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1F, 0x7A, 0x1F)
    # light grey shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    return p


def add_code_block(doc, lines: list[str]):
    for line in lines:
        add_code(doc, line)
    doc.add_paragraph()  # spacer


def add_note(doc, text: str):
    p = doc.add_paragraph()
    run = p.add_run("NOTE: ")
    run.bold = True
    run.font.color.rgb = RGBColor(0xB8, 0x60, 0x00)
    p.add_run(text)
    return p


def add_bullet(doc, text: str, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    return p


def add_numbered(doc, text: str):
    return doc.add_paragraph(text, style="List Number")


# ── main ─────────────────────────────────────────────────────────────────────

def build():
    doc = Document()

    # Title
    title = doc.add_heading("DocuMind AI — GCP Deployment Guide", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Terraform-based production deployment on Google Cloud Platform")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    doc.add_paragraph()

    # ── 1. Overview ───────────────────────────────────────────────────────────
    add_heading(doc, "1. Overview")
    add_para(doc, (
        "This guide walks you through a complete, end-to-end deployment of DocuMind AI "
        "on Google Cloud Platform using Terraform. After following these steps you will have:"
    ))
    add_bullet(doc, "A FastAPI backend running on Cloud Run (documind-api)")
    add_bullet(doc, "A Streamlit WebUI running on Cloud Run (documind-webui)")
    add_bullet(doc, "A GCS bucket for uploaded source documents")
    add_bullet(doc, "All secrets stored in Secret Manager (OpenAI key, Qdrant credentials)")
    add_bullet(doc, "A Docker image repository in Artifact Registry")
    add_bullet(doc, "Least-privilege IAM via a dedicated service account")
    doc.add_paragraph()

    # ── 2. Prerequisites ──────────────────────────────────────────────────────
    add_heading(doc, "2. Prerequisites")
    add_para(doc, "Install the following tools before proceeding:", bold=True)

    tools = [
        ("Terraform >= 1.5", "https://developer.hashicorp.com/terraform/downloads"),
        ("Google Cloud SDK (gcloud)", "https://cloud.google.com/sdk/docs/install"),
        ("Docker Desktop", "https://www.docker.com/products/docker-desktop/"),
        ("Python 3.11", "https://www.python.org/downloads/"),
    ]
    for name, url in tools:
        add_bullet(doc, f"{name}  —  {url}")
    doc.add_paragraph()

    add_para(doc, "You also need:", bold=True)
    add_bullet(doc, "A GCP project with billing enabled")
    add_bullet(doc, "An OpenAI API key (platform.openai.com)")
    add_bullet(doc, "A Qdrant Cloud cluster URL and API key (cloud.qdrant.io) — or self-hosted Qdrant")
    doc.add_paragraph()

    # ── 3. GCP Authentication ─────────────────────────────────────────────────
    add_heading(doc, "3. GCP Authentication")
    add_para(doc, "Authenticate your local machine with GCP Application Default Credentials:")
    add_code_block(doc, [
        "gcloud auth login",
        "gcloud auth application-default login",
        "gcloud config set project YOUR_PROJECT_ID",
    ])
    add_note(doc, (
        "Replace YOUR_PROJECT_ID with your actual GCP project ID in every command below. "
        "You can find it in the GCP Console under Project Info."
    ))
    doc.add_paragraph()

    # ── 4. Enable APIs ────────────────────────────────────────────────────────
    add_heading(doc, "4. Enable Required GCP APIs")
    add_para(doc, (
        "Terraform will enable all required APIs automatically via apis.tf. "
        "If you prefer to enable them manually first, run:"
    ))
    add_code_block(doc, [
        "gcloud services enable \\",
        "  run.googleapis.com \\",
        "  storage.googleapis.com \\",
        "  secretmanager.googleapis.com \\",
        "  artifactregistry.googleapis.com \\",
        "  cloudbuild.googleapis.com \\",
        "  iam.googleapis.com \\",
        "  cloudresourcemanager.googleapis.com \\",
        "  logging.googleapis.com \\",
        "  monitoring.googleapis.com",
    ])

    # ── 5. Terraform Setup ────────────────────────────────────────────────────
    add_heading(doc, "5. Terraform Setup")

    add_heading(doc, "5.1 Create a tfvars file", level=2)
    add_para(doc, "Copy the example and fill in your values:")
    add_code_block(doc, [
        "cd terraform",
        "cp terraform.tfvars.example terraform.tfvars",
    ])
    add_para(doc, "Edit terraform.tfvars with your actual values:")
    add_code_block(doc, [
        'project_id      = "my-gcp-project-id"',
        'gcs_bucket_docs = "my-gcp-project-id-documind-docs"',
        "",
        "# Leave api_image and webui_image blank for now — fill after Step 7",
        'api_image   = ""',
        'webui_image = ""',
        "",
        'region          = "us-central1"',
        'collection_name = "documindai"',
        'environment     = "production"',
    ])

    add_heading(doc, "5.2 (Optional) Remote State Backend", level=2)
    add_para(doc, (
        "For teams or CI/CD, store Terraform state in GCS so it is shared and locked. "
        "Skip this step for solo development."
    ))
    add_code_block(doc, [
        "# Create a state bucket (one-time)",
        "gsutil mb -p YOUR_PROJECT_ID -l us-central1 gs://YOUR_PROJECT_ID-tf-state",
    ])
    add_para(doc, "Then uncomment the backend block in terraform/versions.tf:")
    add_code_block(doc, [
        'backend "gcs" {',
        '  bucket = "YOUR_PROJECT_ID-tf-state"',
        '  prefix = "documind-ai/state"',
        "}",
    ])
    add_para(doc, "Migrate state:")
    add_code_block(doc, ["terraform init -migrate-state"])

    add_heading(doc, "5.3 First Apply (infrastructure only)", level=2)
    add_para(doc, (
        "Run Terraform to create all infrastructure except the Cloud Run services "
        "(you need the Docker image first). "
        "Use -target to apply only the non-Cloud-Run resources:"
    ))
    add_code_block(doc, [
        "terraform init",
        "terraform plan",
        "terraform apply \\",
        "  -target=google_project_service.required_apis \\",
        "  -target=google_service_account.documind_api \\",
        "  -target=google_storage_bucket.docs \\",
        "  -target=google_artifact_registry_repository.documind \\",
        "  -target=google_secret_manager_secret.openai_api_key \\",
        "  -target=google_secret_manager_secret.qdrant_url \\",
        "  -target=google_secret_manager_secret.qdrant_api_key",
    ])

    # ── 6. Add Secrets ────────────────────────────────────────────────────────
    add_heading(doc, "6. Add Secret Values")
    add_para(doc, (
        "Terraform creates the Secret Manager containers but does NOT store the secret values "
        "(to avoid them appearing in state files). Add values with gcloud:"
    ))
    add_code_block(doc, [
        "# OpenAI API key",
        'echo -n "sk-YOUR_OPENAI_KEY" | \\',
        "  gcloud secrets versions add documind-openai-api-key --data-file=-",
        "",
        "# Qdrant Cloud URL (e.g. https://abc123.us-east4-0.gcp.cloud.qdrant.io:6333)",
        'echo -n "https://YOUR_QDRANT_URL" | \\',
        "  gcloud secrets versions add documind-qdrant-url --data-file=-",
        "",
        "# Qdrant API key",
        'echo -n "YOUR_QDRANT_API_KEY" | \\',
        "  gcloud secrets versions add documind-qdrant-api-key --data-file=-",
    ])
    add_note(doc, (
        "If using a self-hosted Qdrant (e.g. the Docker Compose setup), "
        "set the Qdrant URL to the internal Cloud Run service URL after deploying it, "
        "or use a public URL if Qdrant is exposed externally."
    ))
    doc.add_paragraph()

    # ── 7. Build & Push Docker Image ──────────────────────────────────────────
    add_heading(doc, "7. Build and Push the Docker Image")

    add_heading(doc, "7.1 Configure Docker for Artifact Registry", level=2)
    add_code_block(doc, [
        "gcloud auth configure-docker us-central1-docker.pkg.dev",
    ])

    add_heading(doc, "7.2 Build the image", level=2)
    add_para(doc, "From the repository root:")
    add_code_block(doc, [
        "# Set variables",
        "PROJECT_ID=your-gcp-project-id",
        "REGION=us-central1",
        "IMAGE_TAG=v1",
        "IMAGE_URI=${REGION}-docker.pkg.dev/${PROJECT_ID}/documind/api:${IMAGE_TAG}",
        "",
        "# Build",
        "docker build -t ${IMAGE_URI} .",
    ])

    add_heading(doc, "7.3 Push the image", level=2)
    add_code_block(doc, [
        "docker push ${IMAGE_URI}",
        "",
        "# Verify",
        "gcloud artifacts docker images list ${REGION}-docker.pkg.dev/${PROJECT_ID}/documind",
    ])

    # ── 8. Final Terraform Apply ──────────────────────────────────────────────
    add_heading(doc, "8. Deploy Cloud Run Services")

    add_heading(doc, "8.1 Update tfvars with the image URI", level=2)
    add_para(doc, "Edit terraform/terraform.tfvars:")
    add_code_block(doc, [
        'api_image   = "us-central1-docker.pkg.dev/YOUR_PROJECT_ID/documind/api:v1"',
        'webui_image = "us-central1-docker.pkg.dev/YOUR_PROJECT_ID/documind/api:v1"',
    ])

    add_heading(doc, "8.2 Apply the full configuration", level=2)
    add_code_block(doc, [
        "terraform plan",
        "terraform apply",
    ])
    add_para(doc, "Terraform will output the deployed URLs:")
    add_code_block(doc, [
        "Outputs:",
        "",
        'api_url    = "https://documind-api-abc123-uc.a.run.app"',
        'webui_url  = "https://documind-webui-abc123-uc.a.run.app"',
        'gcs_bucket_docs = "my-gcp-project-id-documind-docs"',
    ])

    # ── 9. Verify the Deployment ──────────────────────────────────────────────
    add_heading(doc, "9. Verify the Deployment")

    add_heading(doc, "9.1 Health check the API", level=2)
    add_code_block(doc, [
        "curl https://documind-api-HASH-uc.a.run.app/health",
        "# Expected: {\"status\": \"ok\"}",
    ])

    add_heading(doc, "9.2 Test a query", level=2)
    add_code_block(doc, [
        'curl -X POST https://documind-api-HASH-uc.a.run.app/api/v1/query \\',
        '  -H "Content-Type: application/json" \\',
        "  -d '{\"question\": \"What is DocuMind AI?\", \"session_id\": \"test\"}'",
    ])

    add_heading(doc, "9.3 Open the WebUI", level=2)
    add_para(doc, "Navigate to the webui_url output in a browser. You should see the Streamlit chat interface.")
    doc.add_paragraph()

    # ── 10. Upload Documents ──────────────────────────────────────────────────
    add_heading(doc, "10. Upload Documents to GCS")
    add_para(doc, (
        "Upload source documents to the GCS bucket. "
        "The application indexes them on upload via the API."
    ))
    add_code_block(doc, [
        "# Copy local documents to GCS",
        "gsutil cp -r data/raw/* gs://YOUR_PROJECT_ID-documind-docs/raw/",
    ])
    add_para(doc, "Or use the WebUI upload button to upload documents directly through the browser.")
    doc.add_paragraph()

    # ── 11. Terraform File Reference ──────────────────────────────────────────
    add_heading(doc, "11. Terraform File Reference")

    files = [
        ("versions.tf", "Terraform version constraint and optional GCS backend"),
        ("provider.tf", "Google provider configuration"),
        ("variables.tf", "All input variables with descriptions and defaults"),
        ("locals.tf", "Computed locals: Artifact Registry base URL, resource labels"),
        ("apis.tf", "Enables all 9 required GCP APIs"),
        ("iam.tf", "Service account and IAM role bindings"),
        ("secrets.tf", "Secret Manager secret containers (values added manually)"),
        ("gcs.tf", "GCS bucket with versioning and lifecycle rules"),
        ("artifact_registry.tf", "Docker image repository"),
        ("cloud_run_api.tf", "FastAPI Cloud Run service with health probes"),
        ("cloud_run_webui.tf", "Streamlit Cloud Run service with CMD override"),
        ("outputs.tf", "Post-apply outputs: URLs, bucket name, SA email"),
        ("terraform.tfvars.example", "Example variables file — copy and customise"),
    ]
    for fname, desc in files:
        p = doc.add_paragraph()
        p.add_run(f"  {fname}").bold = True
        p.add_run(f"  —  {desc}")
    doc.add_paragraph()

    # ── 12. Scaling & Cost Tuning ─────────────────────────────────────────────
    add_heading(doc, "12. Scaling and Cost Tuning")

    add_para(doc, "Key variables that affect cost and cold-start behaviour:", bold=True)
    add_code_block(doc, [
        "api_min_instances  = 0   # scale-to-zero (free when idle); set to 1 to eliminate cold starts",
        "api_max_instances  = 10  # upper bound on concurrent instances",
        "webui_min_instances = 0",
        "webui_max_instances = 3",
    ])

    add_para(doc, "Resources per container (edit cloud_run_api.tf to change):", bold=True)
    add_code_block(doc, [
        "API:   2 vCPU, 4 GiB RAM",
        "WebUI: 1 vCPU, 2 GiB RAM",
    ])
    doc.add_paragraph()

    # ── 13. CI/CD (optional) ──────────────────────────────────────────────────
    add_heading(doc, "13. CI/CD with Cloud Build (optional)")
    add_para(doc, (
        "A basic Cloud Build trigger can rebuild and redeploy on every push to main. "
        "Create a cloudbuild.yaml at the repo root:"
    ))
    add_code_block(doc, [
        "steps:",
        "  - name: gcr.io/cloud-builders/docker",
        "    args:",
        "      - build",
        "      - -t",
        "      - $_IMAGE_URI",
        "      - .",
        "  - name: gcr.io/cloud-builders/docker",
        "    args: [push, $_IMAGE_URI]",
        "  - name: gcr.io/google.com/cloudsdktool/cloud-sdk",
        "    entrypoint: gcloud",
        "    args:",
        "      - run",
        "      - services",
        "      - update",
        "      - documind-api",
        "      - --image=$_IMAGE_URI",
        "      - --region=$_REGION",
        "substitutions:",
        "  _REGION: us-central1",
        "  _IMAGE_URI: us-central1-docker.pkg.dev/$PROJECT_ID/documind/api:$COMMIT_SHA",
    ])

    # ── 14. Troubleshooting ───────────────────────────────────────────────────
    add_heading(doc, "14. Troubleshooting")

    issues = [
        (
            "terraform apply fails: 'API not enabled'",
            (
                "The APIs take up to 60 seconds to activate after enablement. "
                "Re-run terraform apply — it is idempotent."
            ),
        ),
        (
            "Cloud Run service fails: 'Secret not found'",
            (
                "You created the secret container but haven't added a version yet. "
                "Run the gcloud secrets versions add commands from Step 6."
            ),
        ),
        (
            "docker push: 'denied: Permission denied'",
            (
                "Run: gcloud auth configure-docker us-central1-docker.pkg.dev  "
                "and ensure your account has roles/artifactregistry.writer."
            ),
        ),
        (
            "API returns 500: 'OPENAI_API_KEY not set'",
            (
                "The Cloud Run service couldn't access the secret. Verify the service account "
                "has roles/secretmanager.secretAccessor on the secret (iam.tf handles this)."
            ),
        ),
        (
            "WebUI shows 'Connection refused' to API",
            (
                "The API_BASE_URL env var is set to the Cloud Run API URI automatically by Terraform "
                "(cloud_run_webui.tf). Confirm the API service deployed successfully first."
            ),
        ),
        (
            "Qdrant filter returns 0 results",
            (
                "A payload index on metadata.doc_id is required. "
                "The application creates it automatically on first use. "
                "Re-upload the document if results are empty after first deployment."
            ),
        ),
    ]
    for title, solution in issues:
        p = doc.add_paragraph()
        p.add_run(f"Problem: {title}").bold = True
        doc.add_paragraph(f"Solution: {solution}")
        doc.add_paragraph()

    # ── 15. Teardown ─────────────────────────────────────────────────────────
    add_heading(doc, "15. Teardown")
    add_para(doc, (
        "To destroy all GCP resources created by Terraform "
        "(WARNING: this permanently deletes data):"
    ))
    add_code_block(doc, [
        "terraform destroy",
    ])
    add_note(doc, (
        "The GCS bucket has force_destroy = false. "
        "Empty it manually before destroy: gsutil rm -r gs://YOUR_BUCKET_NAME/**"
    ))
    doc.add_paragraph()

    # ── 16. Quick-reference Checklist ────────────────────────────────────────
    add_heading(doc, "16. Deployment Checklist")

    steps = [
        "Install Terraform, gcloud, Docker",
        "gcloud auth login && gcloud auth application-default login",
        "gcloud config set project YOUR_PROJECT_ID",
        "cd terraform && cp terraform.tfvars.example terraform.tfvars",
        "Fill in project_id and gcs_bucket_docs in terraform.tfvars",
        "terraform init",
        "terraform apply -target=... (infrastructure without Cloud Run)",
        "Add secret values via gcloud secrets versions add (Step 6)",
        "gcloud auth configure-docker us-central1-docker.pkg.dev",
        "docker build -t IMAGE_URI . && docker push IMAGE_URI",
        "Update api_image and webui_image in terraform.tfvars",
        "terraform apply (deploys Cloud Run services)",
        "curl https://API_URL/health — confirm 200 OK",
        "Open WebUI URL in browser, upload a document, ask a question",
    ]
    for step in steps:
        add_numbered(doc, step)
    doc.add_paragraph()

    # footer
    p = doc.add_paragraph()
    p.add_run("DocuMind AI | Developer: Anu Sasidharan | Last updated: May 2025").italic = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(OUT))
    print(f"Written: {OUT}")


if __name__ == "__main__":
    try:
        build()
    except ImportError as e:
        print(f"Missing dependency: {e}")
        print("Install with:  pip install python-docx")
        sys.exit(1)
