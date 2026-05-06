"""
Generate docs/InterviewQA.docx — interview Q&A plus architecture diagrams (ASCII + Mermaid text).

  pip install python-docx
  python scripts/build_interview_qa_docx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "InterviewQA.docx"


def add_title(doc: Document, text: str) -> None:
    h = doc.add_heading(text, 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def add_code(doc: Document, text: str) -> None:
    """Monospace block (diagrams, Mermaid, code)."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.rstrip("\n"))
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def qa(doc: Document, question: str, answer_paragraphs: list[str]) -> None:
    p = doc.add_paragraph()
    r = p.add_run("Q: ")
    r.bold = True
    p.add_run(question)
    for para in answer_paragraphs:
        ap = doc.add_paragraph()
        ar = ap.add_run("A: ")
        ar.bold = True
        ap.add_run(para)


def main() -> None:
    try:
        import docx  # noqa: F401
    except ImportError:
        print("Install: python -m pip install python-docx", file=sys.stderr)
        sys.exit(1)

    DOCS.mkdir(parents=True, exist_ok=True)
    doc = Document()

    add_title(doc, "DocuMind AI")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st = subtitle.add_run("Interview Questions & Answers\nArchitecture reference (study guide)")
    st.bold = False
    st.font.size = Pt(12)

    add_p(
        doc,
        "Generated from the DocuMind AI codebase: FastAPI, Streamlit, LangChain, OpenAI embeddings + GPT-4o, "
        "Qdrant dense retrieval, Hugging Face cross-encoder reranking (langchain-classic), Pydantic settings.",
    )
    add_p(
        doc,
        "How to use: Word renders monospace blocks below as diagrams. Copy 'Mermaid' blocks into "
        "mermaid.live or a Markdown viewer if you prefer rendered graphics for slides.",
    )

    add_h(doc, "Reference diagram A — Logical architecture (ASCII)", 1)
    add_code(
        doc,
        """
                    +-----------+       +---------------+
 End user -------> | Streamlit |       |  FastAPI      | <----- API clients
                    +-----+-----+       +-------+-------+
                          |                     |
                          |  import query()     |  POST /api/v1/query
                          v                     v
                    +-----------------------------------+
                    |  src/rag_pipeline.py               |
                    |  RunnableWithMessageHistory        |
                    |  -> retriever (reranked) -> LLM    |
                    +--+-------------+----------+--------+
                       |             |          |
                       v             v          v
                 +----------+ +----------+ +---------+
                 | Qdrant   | | Cross-   | | OpenAI  |
                 | vectors  | | encoder  | | embed + |
                 |          | | rerank   | | chat    |
                 +----------+ +----------+ +---------+
""",
    )

    add_h(doc, "Reference diagram B — Query sequence (ASCII)", 1)
    add_code(
        doc,
        """
  User question + session_id
       |
       v
  [Embed query via OpenAI] ----> Qdrant Top-K similarity
       |
       v
  [Cross-encoder rerank] ------> Top-N passages
       |
       v
  [Build prompt: system + context + chat_history + question]
       |
       v
  [GPT-4o] ------> Markdown answer
""",
    )

    add_h(doc, "Section 1 — Product & problem framing", 1)

    qa(
        doc,
        "What is DocuMind AI, in one sentence?",
        [
            "It is a retrieval-augmented documentation assistant: it searches your indexed technical docs in Qdrant, "
            "reranks the best passages, then asks GPT-4o to answer strictly using that context, with optional per-session chat memory.",
        ],
    )

    qa(
        doc,
        "What business problem does it solve?",
        [
            "Developers and support teams waste time hunting through long documentation. DocuMind reduces time-to-answer "
            "and improves accuracy by grounding responses in the actual corpus instead of relying on the model’s parametric memory.",
        ],
    )

    qa(
        doc,
        "Who are the users and how do they access the system?",
        [
            "End users interact through Streamlit for demos or through the FastAPI REST API for integration. Operators run "
            "indexing scripts (or future batch jobs) to load markdown/text into Qdrant. In production you would add authentication.",
        ],
    )

    add_h(doc, "Section 2 — RAG fundamentals", 1)

    qa(
        doc,
        "Explain RAG in simple terms.",
        [
            "Retrieval-Augmented Generation first retrieves relevant text chunks from a knowledge base, then passes them to a "
            "language model as context. The model generates an answer conditioned on that context, which reduces fabricated "
            "details compared to ‘the model alone’ when facts must match your documents.",
        ],
    )

    qa(
        doc,
        "Why not just use ChatGPT on the whole document set?",
        [
            "Context windows and cost: you cannot fit all docs in one prompt. Even if you could, models attend poorly to "
            "huge contexts. Retrieval narrows down to dozens of passages. It also lets you cite sources and refresh knowledge "
            "by re-indexing without retraining.",
        ],
    )

    qa(
        doc,
        "What embedding model does DocuMind use and why?",
        [
            "Settings default to OpenAI text-embedding-3-small. It is a good balance of quality, latency, and cost for "
            "semantic search; vectors are consumed by Qdrant with cosine distance. You could swap dimensions/model in src/config.py.",
        ],
    )

    qa(
        doc,
        "Dense retrieval vs keyword (BM25) search — compare them.",
        [
            "Dense retrieval matches meaning (synonyms, paraphrases). BM25 matches exact or lexical overlap. Dense can miss precise "
            "identifiers; BM25 can miss conceptual questions. Many production systems use hybrid retrieval (ensemble) for both.",
        ],
    )

    qa(
        doc,
        "Does this codebase implement hybrid BM25 + dense search?",
        [
            "Not yet. The retriever is Qdrant similarity (dense) only, followed by cross-encoder reranking. A natural extension is "
            "LangChain’s EnsembleRetriever with BM25Retriever from the same chunk set and tuned weights (e.g. 0.4 BM25, 0.6 dense).",
        ],
    )

    qa(
        doc,
        "Why do you have a reranker after vector search?",
        [
            "Vector search optimizes approximate similarity at scale; ordering is imperfect. A cross-encoder scores each "
            "(query, passage) pair jointly, which is slower but more accurate for the top results. You retrieve K=8 then "
            "rerank to top_n=5 before building the LLM prompt, improving grounding and reducing noise.",
        ],
    )

    qa(
        doc,
        "What reranker is used and could you use Cohere instead?",
        [
            "The code uses HuggingFaceCrossEncoder with a configurable model (default cross-encoder/ms-marco-MiniLM-L-6-v2) via "
            "CrossEncoderReranker in langchain_classic. You could swap to Cohere Rerank v3 using LangChain’s Cohere integrations "
            "and treat it as a managed API trade-off (latency/cost vs ops).",
        ],
    )

    add_h(doc, "Reference diagram — Mermaid: system context (paste into mermaid.live)", 2)
    add_code(
        doc,
        """
flowchart TB
    U[Users] --> ST[Streamlit]
    U --> API[FastAPI]
    ST --> CORE[rag_pipeline.query]
    API --> CORE
    CORE --> QD[Qdrant]
    CORE --> OAI[OpenAI]
""",
    )

    add_h(doc, "Section 3 — Architecture & components", 1)

    qa(
        doc,
        "Describe the runtime architecture at container level.",
        [
            "The Dockerfile runs Uvicorn on api.main:app (port 8000). It bundles src/, api/, pipelines/, scripts/. Qdrant is not "
            "inside that image—it is external (Docker Compose locally, Qdrant Cloud or another service in GCP). Streamlit runs as a separate process for demos.",
        ],
    )

    qa(
        doc,
        "What is in src/config.py?",
        [
            "A Pydantic BaseSettings class with typed fields: OPENAI_API_KEY, QDRANT_URL, COLLECTION_NAME, embedding and LLM model "
            "names, chunk_size/overlap, top_k, rerank_top_n, CROSS_ENCODER_MODEL, optional GCP vars. get_settings() is lru_cached "
            "for efficiency.",
        ],
    )

    qa(
        doc,
        "What does src/retriever.py encapsulate?",
        [
            "Qdrant client construction, QdrantVectorStore with cosine distance, shared OpenAIEmbeddings, and a similarity retriever "
            "with search_kwargs.k = top_k. This is stage-one recall.",
        ],
    )

    qa(
        doc,
        "What does src/reranker.py do?",
        [
            "It wraps the base retriever in ContextualCompressionRetriever with CrossEncoderReranker, returning fewer, better-ordered "
            "documents before the prompt is built.",
        ],
    )

    qa(
        doc,
        "How is the LangChain chain structured in rag_pipeline.py?",
        [
            "build_rag_chain uses RunnableParallel: one branch pipes question -> retriever -> format_docs into context string; parallel "
            "branches carry question and chat_history. That dict feeds a ChatPromptTemplate (system with {context}, "
            "MessagesPlaceholder for history, human question), then ChatOpenAI, then StrOutputParser.",
        ],
    )

    qa(
        doc,
        "How does conversational memory work?",
        [
            "RunnableWithMessageHistory wraps the chain. get_session_history maps session_id to InMemoryChatMessageHistory. Input key is "
            "question; history key is chat_history. For production you would persist history in Redis/Firestore and bound length.",
        ],
    )

    qa(
        doc,
        "Why cache _rag_with_history_default?",
        [
            "Lazy singleton avoids reloading the cross-encoder Hugging Face model on every request, reducing cold latency after "
            "the first invocation in a warm instance.",
        ],
    )

    qa(
        doc,
        "How does indexing work end-to-end in code?",
        [
            "index_documents loads paths via pipelines/components/loader.py (recursive .md/.txt), splits with RecursiveCharacterTextSplitter "
            "from chunker.py, then vector_store.add_documents in Qdrant (embeddings via LangChain/OpenAI integration). Scripts/index_docs.py is the CLI wrapper.",
        ],
    )

    add_h(doc, "Reference diagram — Mermaid: indexing flow", 2)
    add_code(
        doc,
        """
flowchart LR
    P[paths] --> L[load_text_files]
    L --> SP[RecursiveCharacterTextSplitter]
    SP --> ADD[vector_store.add_documents]
    ADD --> Q[(Qdrant)]
    ADD --> E[OpenAI embeddings]
""",
    )

    qa(
        doc,
        "What are the FastAPI endpoints?",
        [
            "GET /health for liveness-style checks; POST /api/v1/query with JSON { question, session_id } returning { answer, session_id }. Root GET returns pointers to docs.",
        ],
    )

    qa(
        doc,
        "Why import query inside the route handler?",
        [
            "It defers importing the heavy LangChain/OpenAI/Qdrant stack until the endpoint is exercised, slightly improving startup of "
            "the API container in environments where imports matter. Alternative: lifespan hooks or eager init for predictable cold start.",
        ],
    )

    add_h(doc, "Section 4 — Prompt & safety", 1)

    qa(
        doc,
        "How do you instruct the model to stay grounded?",
        [
            "The system prompt says to answer based on provided context, cite source sections, include code examples when relevant, "
            "and explicitly state when information is not in the context. That pattern reduces confident hallucinations on missing facts.",
        ],
    )

    qa(
        doc,
        "What temperature do you use and why?",
        [
            "Settings use temperature 0.1 — low randomness for factual, documentation-style answers where consistency beats creativity.",
        ],
    )

    add_h(doc, "Section 5 — Evaluation & quality", 1)

    qa(
        doc,
        "How would you measure retrieval quality?",
        [
            "Offline: labeled question–chunk relevance sets; metrics like Precision@k, Recall@k, nDCG, MRR on your corpus or a sampled "
            "benchmark. Compare dense-only vs dense+rerank vs future hybrid.",
        ],
    )

    qa(
        doc,
        "How would you measure end-to-end answer quality?",
        [
            "Frameworks such as RAGAS (faithfulness, answer relevance, context precision/recall), human eval rubrics, A/B testing "
            "in production with user feedback thumbs, or LLM-as-judge with safeguards.",
        ],
    )

    qa(
        doc,
        "What would you log in production?",
        [
            "Request id, latency per stage (retrieval, rerank, LLM), token counts, model versions, retrieval scores, chunk ids, errors, "
            "and user/session keys (hashed) for GDPR-aware analytics.",
        ],
    )

    add_h(doc, "Section 6 — GCP, security, ops (Datatonic angle)", 1)

    qa(
        doc,
        "How would you deploy this on Google Cloud?",
        [
            "Build the Dockerfile to Artifact Registry. Run Cloud Run for the FastAPI service with min/max instances tuned for demos. "
            "Store raw docs in GCS (versioned buckets). Secrets (OpenAI, QDRANT_URL) in Secret Manager, injected as env vars. "
            "Use Terraform for repeatable IAM + resources; see GCP_deployment_guide.docx in the repo root.",
        ],
    )

    qa(
        doc,
        "Where does Qdrant run in GCP?",
        [
            "Not in this image. Options: Qdrant Cloud, a second Cloud Run service with persistence trade-offs, GKE StatefulSet, "
            "or managed vector offerings as they fit latency and SLA needs.",
        ],
    )

    qa(
        doc,
        "Failure modes and mitigations?",
        [
            "OpenAI rate limits or outages: exponential backoff, circuit breakers, cached answers for FAQs, degraded mode messaging. "
            "Qdrant unavailable: health checks, retries, failover read replica pattern. Empty index: surface clear 'not indexed yet' messaging.",
        ],
    )

    qa(
        doc,
        "Cost levers?",
        [
            "Reduce LLM tokens by tightening top_n and shortening chunks; smaller LLM for simple queries routing; caching embeddings "
            "for static chunks; hybrid search to shrink bad retrievals; min instances=0 on Cloud Run for dev.",
        ],
    )

    qa(
        doc,
        "Data privacy considerations?",
        [
            "User questions may contain sensitive strings—avoid logging raw PII, redact secrets, define retention for chat history "
            "if persisted, encrypt GCS buckets, restrict invoker IAM on Cloud Run instead of public allUsers where possible.",
        ],
    )

    add_h(doc, "Reference diagram — Mermaid: target GCP deployment", 2)
    add_code(
        doc,
        """
flowchart TB
    U[Users HTTPS] --> CR[Cloud Run documind-api]
    CR --> SM[Secret Manager]
    CR --> GCS[GCS docs bucket]
    CR --> LOG[Cloud Logging]
    CR --> QD[(Qdrant)]
    CR --> OAI[OpenAI APIs]
""",
    )

    add_h(doc, "Section 7 — Scaling & roadmap", 1)

    qa(
        doc,
        "How does this scale to millions of chunks?",
        [
            "Sharding collections, multi-replica Qdrant, quantization and ANN tuning, batch indexing via Vertex pipelines, "
            "asynchronous ingestion, CDN for assets, caching hot queries.",
        ],
    )

    qa(
        doc,
        "What improvements would you add next?",
        [
            "Hybrid BM25+dense retrieval; managed Cohere rerank optional path; citations with chunk IDs and URIs stored in payload; "
            "Cloud Run Jobs for scheduled reindex from GCS; Prometheus metrics; authenticated API Gateway; multilingual embeddings.",
        ],
    )

    qa(
        doc,
        "Limitations of InMemoryChatMessageHistory?",
        [
            "It is process-local and lost on restart; not suited for horizontally scaled replicas without sticky sessions unless "
            "replaced by external conversation store.",
        ],
    )

    add_h(doc, "Section 8 — Code hygiene & testing", 1)

    qa(
        doc,
        "What tests exist?",
        [
            "Tests cover chunk splitter behavior on long documents, file loader parsing, FastAPI health route; full RAG integration test "
            "is marked skipped until env + Qdrant + indexed data are available.",
        ],
    )

    qa(
        doc,
        "Why LangChain instead of raw SDK calls?",
        [
            "Composable runnables (LCEL), retriever abstractions, and integration with multiple vector stores and models speed "
            "iteration. Trade-off: framework overhead and version churn—mitigate by keeping core logic small in rag_pipeline.py.",
        ],
    )

    qa(
        doc,
        "Why does reranker code import from langchain_classic?",
        [
            "In LangChain v1 layouts, contextual compression retrievers and cross-encoder compressors live under the langchain-classic "
            "compatibility package; this project pins that stack so imports stay stable alongside langchain-community cross-encoder adapters.",
        ],
    )

    qa(
        doc,
        "What does operator.itemgetter do in the parallel chain?",
        [
            "It pulls keys (question, chat_history) from the dict that RunnableWithMessageHistory passes into the inner chain, while the "
            "context branch computes document text from question | retriever | format_docs.",
        ],
    )

    qa(
        doc,
        "Do returned answers include explicit source filenames today?",
        [
            "The prompt asks for citations, but loader metadata (e.g. source path) is only as rich as LangChain attaches to Document "
            "metadata. A strong next step is to persist file path and heading in chunk metadata and surface it in format_docs or the prompt.",
        ],
    )

    qa(
        doc,
        "Is re-running indexing idempotent?",
        [
            "add_documents appends vectors; duplicates can appear if you ingest the same path twice without delete/upsert policy. Production "
            "pipelines usually use deterministic ids, versioning, or wipe collection before full reindex.",
        ],
    )

    qa(
        doc,
        "How do you choose chunk_size and overlap?",
        [
            "They trade context vs noise: larger chunks carry more surrounding text but dilute retrieval; overlap preserves sentences split "
            "across boundaries. Tune with offline eval (hit rate / user satisfaction) starting from 512 tokens and ~50 overlap as in Settings.",
        ],
    )

    qa(
        doc,
        "Could you run the cross-encoder on GPU?",
        [
            "Yes for self-hosted inference; sentence-transformers/torch will use CUDA if installed and configured. Cloud Run CPUs are typical "
            "for MS-MARCO Mini models; GPU would move to GKE or specialized serving if latency dominates.",
        ],
    )

    qa(
        doc,
        "How does docker-compose.yml fit local dev?",
        [
            "It runs only the Qdrant container on ports 6333/6334 with a named volume—lightweight substitute for managed vector DB during development.",
        ],
    )

    qa(
        doc,
        "What would you say about fine-tuning vs RAG for this use case?",
        [
            "RAG updates facts by editing documents and re-indexing—fast for evolving docs. Fine-tuning improves style or domain phrasing "
            "but stale knowledge remains a risk unless paired with continual training. Hybrid: RAG for facts, light fine-tune for tone if needed.",
        ],
    )

    qa(
        doc,
        "How would you integrate GCS ingestion with existing code?",
        [
            "Download objects to tmp or stream to memory in a new loader component, reuse chunker/embed pipeline, optionally trigger from "
            "Cloud Storage notifications via Cloud Functions/Run Jobs. Keep secrets and service accounts scoped to objectViewer on the docs bucket.",
        ],
    )

    add_h(doc, "Section 9 — Quick flashcard answers", 1)
    flash = [
        ("top_k vs rerank_top_n?", "top_k is how many candidates Qdrant returns; rerank_top_n is how many survive cross-encoder scoring for the prompt."),
        ("Distance metric?", "Cosine similarity on embeddings in QdrantVectorStore configuration."),
        ("Collection name?", "Configurable via COLLECTION_NAME env defaulting to docs-index."),
        ("Languages supported?", "Only as far as embeddings/LLM support; no explicit language routing in code yet."),
        ("PDF support?", "Not in loader.py today—would add parsers or unstructured pipeline."),
    ]
    for qtext, atext in flash:
        qa(doc, qtext, [atext])

    add_h(doc, "Appendix — Full Mermaid: query sequence", 1)
    add_code(
        doc,
        """
sequenceDiagram
    participant U as User
    participant S as Streamlit or API
    participant R as rag_pipeline
    participant Q as Qdrant retriever
    participant C as Cross-encoder rerank
    participant L as GPT-4o
    U->>S: question + session_id
    S->>R: query()
    R->>Q: similarity search (embed query)
    Q-->>R: Top-K chunks
    R->>C: rerank
    C-->>R: Top-N chunks
    R->>L: prompt with context + history
    L-->>U: answer
""",
    )

    add_p(doc, "End of document — regenerate with: python scripts/build_interview_qa_docx.py")

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
